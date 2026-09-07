import os
import logging
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("excel_exporter")

# Resolve project root (parent directory of backend)
BACKEND_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(BACKEND_DIR, ".."))
DEFAULT_EXCEL_FILENAME = "shieldscope_scan_report.xlsx"

# Header configuration
HEADERS = [
    "URL",
    "Category",
    "Score",
    "Scanned At",
    "Exposed Paths",
    "Header Issues",
    "CORS Issues"
]

COLUMN_WIDTHS = {
    "A": 36,  # URL
    "B": 16,  # Category
    "C": 10,  # Score
    "D": 22,  # Scanned At
    "E": 35,  # Exposed Paths
    "F": 35,  # Header Issues
    "G": 35   # CORS Issues
}

# Styles
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

DATA_FONT = Font(name="Arial", size=10)
DATA_ALIGN_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
DATA_ALIGN_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB")
)


def _format_exposed_paths(exposed_paths_data) -> str:
    """Format exposed paths into line-separated items."""
    if not exposed_paths_data:
        return "None (Secure)"

    if isinstance(exposed_paths_data, dict):
        exposed_list = exposed_paths_data.get("exposed_paths", [])
    elif isinstance(exposed_paths_data, list):
        exposed_list = exposed_paths_data
    else:
        exposed_list = []

    if not exposed_list:
        return "None (Secure)"

    lines = []
    for item in exposed_list:
        if isinstance(item, dict):
            path = item.get("path") or item.get("url") or "Unknown"
            severity = item.get("severity") or "INFO"
            lines.append(f"[{severity}] {path}")
        else:
            lines.append(str(item))

    return "\n".join(lines) if lines else "None (Secure)"


def _format_header_issues(headers_data) -> str:
    """Format missing security headers into line-separated items."""
    if not headers_data or not isinstance(headers_data, dict):
        return "None"

    missing = headers_data.get("missing_headers", [])
    if not missing:
        return "None (All Secure)"

    return "\n".join(f"Missing: {h}" for h in missing)


def _format_cors_issues(cors_data) -> str:
    """Format CORS findings into line-separated items."""
    if not cors_data or not isinstance(cors_data, dict):
        return "None"

    findings = cors_data.get("findings", [])
    if not findings:
        return "None (Secure)"

    issues = []
    for f in findings:
        if isinstance(f, dict):
            sev = f.get("severity", "INFO")
            title = f.get("issue") or f.get("title") or "CORS issue"
            # Filter out informational pass messages
            if sev == "INFO" and "secure" in title.lower():
                continue
            issues.append(f"[{sev}] {title}")
        else:
            issues.append(str(f))

    return "\n".join(issues) if issues else "None (Secure)"


def append_scan_to_excel(scan_data: dict, filepath: str = DEFAULT_EXCEL_FILENAME) -> bool:
    """
    Appends a single scan result row to the Excel report workbook.
    
    If the file does not exist, it creates a new workbook with styled headers,
    frozen top row, and preset column widths.
    
    Catches file lock / PermissionError gracefully without throwing exceptions.
    """
    try:
        if not scan_data:
            logger.warning("Empty scan_data passed to append_scan_to_excel")
            return False

        # Resolve full path
        if not os.path.isabs(filepath):
            full_path = os.path.join(PROJECT_ROOT, filepath)
        else:
            full_path = filepath

        # Extract row fields
        url = scan_data.get("url") or scan_data.get("website") or "N/A"
        category = scan_data.get("category") or "Other"

        # Score extraction
        score = scan_data.get("score")
        if score is None:
            score = scan_data.get("summary", {}).get("security_score", "N/A")

        # Timestamp extraction
        scanned_at = scan_data.get("createdAt")
        if not scanned_at:
            scanned_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        else:
            try:
                # If ISO format, convert to cleaner display format
                dt = datetime.fromisoformat(scanned_at.replace("Z", "+00:00"))
                scanned_at = dt.strftime("%Y-%m-%d %H:%M:%S UTC")
            except Exception:
                pass

        scans_dict = scan_data.get("scans") or {}
        exposed_paths_str = _format_exposed_paths(scans_dict.get("exposed_paths") or scan_data.get("exposed_paths"))
        header_issues_str = _format_header_issues(scans_dict.get("headers") or scan_data.get("headers"))
        cors_issues_str = _format_cors_issues(scans_dict.get("cors") or scan_data.get("cors"))

        row_values = [
            url,
            category,
            score,
            scanned_at,
            exposed_paths_str,
            header_issues_str,
            cors_issues_str
        ]

        # Load or create workbook
        file_exists = os.path.exists(full_path)
        if file_exists:
            wb = openpyxl.load_workbook(full_path)
            ws = wb.active
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Scan Report"

            # Setup Header Row
            ws.append(HEADERS)
            ws.row_dimensions[1].height = 28
            ws.freeze_panes = "A2"

            for col_idx, header in enumerate(HEADERS, 1):
                col_letter = get_column_letter(col_idx)
                cell = ws.cell(row=1, column=col_idx)
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = HEADER_ALIGNMENT
                ws.column_dimensions[col_letter].width = COLUMN_WIDTHS.get(col_letter, 25)

        # Append Data Row
        ws.append(row_values)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = max(24, 18 * max(
            exposed_paths_str.count("\n") + 1,
            header_issues_str.count("\n") + 1,
            cors_issues_str.count("\n") + 1
        ))

        # Apply formatting to data cells
        for col_idx in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            # Center-align Category, Score, and Date; Left-align others
            if col_idx in (2, 3, 4):
                cell.alignment = DATA_ALIGN_CENTER
            else:
                cell.alignment = DATA_ALIGN_LEFT

        # Save workbook
        wb.save(full_path)
        logger.info(f"Successfully appended scan result for {url} to Excel: {full_path}")
        return True

    except PermissionError as pe:
        logger.warning(
            f"Excel report file is currently open in another application (locked): {pe}. "
            f"Skipping Excel write for this scan."
        )
        print(f"[WARNING] Excel file is locked/open in another program. Skipping Excel append: {pe}")
        return False
    except Exception as e:
        logger.error(f"Error appending scan to Excel report: {e}", exc_info=True)
        print(f"[ERROR] Could not append scan to Excel report: {e}")
        return False
